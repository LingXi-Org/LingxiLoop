"""Development-only generator for entirely synthetic issue #424 fixtures."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
import io
import subprocess
import sys
import tempfile
import zipfile

root = Path(__file__).resolve().parent
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[5])
slide.shapes.title.text = 'Synthetic material 合成课程'
slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(2)).text_frame.text = 'Explain why 2 + 3 = 5. 独立说明加法。'
slide.notes_slide.notes_text_frame.text = 'Synthetic speaker notes 说明推理。'
table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(6), Inches(1)).table
for cell, text in zip([table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)], ['Input 输入', 'Output 输出', '2 + 3', '5']):
    cell.text = text
slide = deck.slides.add_slide(deck.slide_layouts[5])
slide.shapes.title.text = 'Second slide 第二页'
slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(2)).text_frame.text = 'Use a different example. 举出不同例子。'
ids = deck.slides._sldIdLst
ids.insert(0, ids[-1])
deck.save(root / 'material.pptx')

deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[6])
picture = io.BytesIO()
Image.new('RGB', (8, 8), 'white').save(picture, format='PNG')
picture.seek(0)
slide.shapes.add_picture(picture, Inches(1), Inches(1))
deck.save(root / 'image-only.pptx')
with zipfile.ZipFile(root / 'expanded-limit.pptx', 'w', compression=zipfile.ZIP_DEFLATED) as archive:
    archive.writestr('ppt/slides/slide1.xml', b' ' * (32 * 1024 * 1024 + 1))

if '--legacy' in sys.argv:
    with tempfile.TemporaryDirectory(prefix='lingxi-synthetic-ppt-') as profile:
        subprocess.run(['soffice', '-env:UserInstallation=' + Path(profile).as_uri(), '--headless', '--convert-to', 'ppt', '--outdir', str(root), str(root / 'material.pptx')], check=True)
